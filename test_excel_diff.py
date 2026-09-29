"""Basit testler: python -m pytest  (veya python test_excel_diff.py)"""
import os
import tempfile

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from excel_diff import compare_workbooks, main


def _make(path, modify=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Veri"
    ws["A1"] = "Baslik"
    ws["A2"] = 10
    ws["A3"] = "=A2*2"
    ws["B1"] = "Font"
    wb.create_sheet("Ortak")["A1"] = "ayni"
    if not modify:
        wb.create_sheet("Silinecek")
    else:
        ws["A2"] = 11                                     # deger farki
        ws["B1"].font = Font(name="Arial", bold=True)     # font farki
        ws["A1"].fill = PatternFill("solid", fgColor="FFFF00")  # dolgu farki
        ws["C5"].comment = Comment("not", "yazar")        # yorum farki
        ws.merge_cells("D1:E1")                           # birlestirme farki
        ws.column_dimensions["A"].width = 30              # sutun genisligi
        wb.create_sheet("Yeni")                           # yeni sheet
    wb.save(path)


def test_diffs():
    with tempfile.TemporaryDirectory() as d:
        p1, p2 = os.path.join(d, "a.xlsx"), os.path.join(d, "b.xlsx")
        _make(p1)
        _make(p2, modify=True)

        diffs, _, _ = compare_workbooks(p1, p2)
        found = {(d.sheet, d.location, d.category, d.prop) for d in diffs}

        assert ("Veri", "A2", "Deger", "Icerik") in found
        assert ("Veri", "B1", "Font", "Ad") in found
        assert ("Veri", "B1", "Font", "Kalin") in found
        assert ("Veri", "A1", "Dolgu", "Tip") in found
        assert ("Veri", "C5", "Yorum", "Metin") in found
        assert ("Veri", "D1:E1", "Birlestirme", "Birlestirilmis hucre") in found
        assert ("Veri", "Sutun A", "Sutun", "Genislik") in found
        assert ("Silinecek", "-", "Sheet", "Varlik") in found
        assert ("Yeni", "-", "Sheet", "Varlik") in found
        assert not any(d.sheet == "Ortak" for d in diffs)

        out = os.path.join(d, "rapor.xlsx")
        assert main([p1, p2, "-o", out, "-q"]) == 1
        assert os.path.isfile(out)
        assert main([p1, p1, "-q"]) == 0


if __name__ == "__main__":
    test_diffs()
    print("OK")
