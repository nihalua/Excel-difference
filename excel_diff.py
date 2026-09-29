#!/usr/bin/env python3
"""
excel_diff.py - Iki Excel dosyasini (xlsx/xlsm) karsilastirir ve farklari raporlar.

Karsilastirilanlar:
  * Sheet'ler: eklenen / silinen sheet'ler, sheet sirasi, gorunurluk, sekme rengi
  * Hucre degerleri ve formulleri
  * Hucre formatlari: font, dolgu (fill), kenarlik (border), hizalama,
    sayi formati, koruma (protection)
  * Hucre yorumlari ve hyperlink'ler
  * Birlestirilmis hucreler, sutun genislikleri, satir yukseklikleri,
    gizli satir/sutunlar, dondurulmus bolmeler (freeze panes)

Kullanim:
  python excel_diff.py eski.xlsx yeni.xlsx
  python excel_diff.py eski.xlsx yeni.xlsx -o fark_raporu.xlsx
  python excel_diff.py eski.xlsx yeni.xlsx -o fark_raporu.csv

Cikis kodu: 0 = fark yok, 1 = fark var, 2 = hata
"""

import argparse
import csv
import os
import sys
from dataclasses import dataclass

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl gerekli: pip install openpyxl")


@dataclass
class Diff:
    sheet: str
    location: str
    category: str
    prop: str
    old: object
    new: object


# --------------------------------------------------------------------------- #
# Yardimci fonksiyonlar
# --------------------------------------------------------------------------- #
def color_str(color):
    """openpyxl Color nesnesini okunabilir bir metne cevirir."""
    if color is None:
        return None
    ctype = getattr(color, "type", None)
    if ctype == "rgb":
        val = f"rgb:{color.rgb}"
    elif ctype == "theme":
        val = f"theme:{color.theme}"
    elif ctype == "indexed":
        val = f"indexed:{color.indexed}"
    elif ctype == "auto":
        val = "auto"
    else:
        val = str(getattr(color, "rgb", color))
    tint = getattr(color, "tint", 0) or 0
    if tint:
        val += f" tint:{round(tint, 4)}"
    return val


def cell_style(cell):
    """Hucrenin tum format ozelliklerini duz bir sozluk olarak dondurur."""
    f = cell.font
    fill = cell.fill
    b = cell.border
    a = cell.alignment
    p = cell.protection

    style = {
        ("Font", "Ad"): f.name,
        ("Font", "Boyut"): f.size,
        ("Font", "Kalin"): bool(f.bold),
        ("Font", "Italik"): bool(f.italic),
        ("Font", "Alti cizili"): f.underline,
        ("Font", "Ustu cizili"): bool(f.strike),
        ("Font", "Renk"): color_str(f.color),
        ("Font", "Dikey hizalama"): f.vertAlign,
        ("Font", "Anahat"): bool(f.outline),
        ("Font", "Golge"): bool(f.shadow),
        ("Dolgu", "Tip"): getattr(fill, "fill_type", None) or getattr(fill, "type", None),
        ("Dolgu", "On plan rengi"): color_str(getattr(fill, "fgColor", None)),
        ("Dolgu", "Arka plan rengi"): color_str(getattr(fill, "bgColor", None)),
        ("Hizalama", "Yatay"): a.horizontal,
        ("Hizalama", "Dikey"): a.vertical,
        ("Hizalama", "Metni kaydir"): bool(a.wrap_text),
        ("Hizalama", "Sigdirmak icin daralt"): bool(a.shrink_to_fit),
        ("Hizalama", "Girinti"): a.indent,
        ("Hizalama", "Metin donusu"): a.text_rotation,
        ("Sayi formati", "Format"): cell.number_format,
        ("Koruma", "Kilitli"): p.locked,
        ("Koruma", "Gizli"): p.hidden,
    }
    # Gradient dolgu ise ek bilgi
    if fill.__class__.__name__ == "GradientFill":
        style[("Dolgu", "Gradient")] = repr(fill)

    for side_name, tr in (("left", "Sol"), ("right", "Sag"), ("top", "Ust"),
                          ("bottom", "Alt"), ("diagonal", "Capraz")):
        side = getattr(b, side_name, None)
        style[("Kenarlik", f"{tr} stil")] = side.style if side else None
        style[("Kenarlik", f"{tr} renk")] = color_str(side.color) if side else None
    style[("Kenarlik", "Capraz yukari")] = bool(b.diagonalUp)
    style[("Kenarlik", "Capraz asagi")] = bool(b.diagonalDown)
    return style


def fmt(v):
    """Rapor icin degeri metne cevirir."""
    if v is None:
        return "(bos)"
    return str(v)


def load(path, data_only):
    keep_vba = path.lower().endswith(".xlsm")
    return openpyxl.load_workbook(path, data_only=data_only, keep_vba=keep_vba)


# --------------------------------------------------------------------------- #
# Karsilastirma
# --------------------------------------------------------------------------- #
def compare_sheet_props(name, ws1, ws2, diffs):
    def add(loc, cat, prop, o, n):
        if o != n:
            diffs.append(Diff(name, loc, cat, prop, o, n))

    add("-", "Sheet", "Gorunurluk", ws1.sheet_state, ws2.sheet_state)
    add("-", "Sheet", "Sekme rengi",
        color_str(ws1.sheet_properties.tabColor), color_str(ws2.sheet_properties.tabColor))
    add("-", "Sheet", "Dondurulmus bolme", ws1.freeze_panes, ws2.freeze_panes)
    add("-", "Sheet", "Boyut (dimensions)", ws1.dimensions, ws2.dimensions)
    add("-", "Sheet", "Otomatik filtre", ws1.auto_filter.ref, ws2.auto_filter.ref)
    add("-", "Sheet", "Kosullu bicimlendirme sayisi",
        len(list(ws1.conditional_formatting)), len(list(ws2.conditional_formatting)))
    add("-", "Sheet", "Veri dogrulama sayisi",
        len(ws1.data_validations.dataValidation), len(ws2.data_validations.dataValidation))
    add("-", "Sheet", "Sheet korumasi", bool(ws1.protection.sheet), bool(ws2.protection.sheet))

    # Birlestirilmis hucreler
    m1 = {str(r) for r in ws1.merged_cells.ranges}
    m2 = {str(r) for r in ws2.merged_cells.ranges}
    for r in sorted(m1 - m2):
        diffs.append(Diff(name, r, "Birlestirme", "Birlestirilmis hucre", "var", "yok"))
    for r in sorted(m2 - m1):
        diffs.append(Diff(name, r, "Birlestirme", "Birlestirilmis hucre", "yok", "var"))

    # Sutun genislikleri / gizli sutunlar
    cols = set(ws1.column_dimensions.keys()) | set(ws2.column_dimensions.keys())
    for col in sorted(cols, key=lambda c: (len(c), c)):
        d1 = ws1.column_dimensions[col]
        d2 = ws2.column_dimensions[col]
        w1 = d1.width if d1.customWidth else None
        w2 = d2.width if d2.customWidth else None
        add(f"Sutun {col}", "Sutun", "Genislik", w1, w2)
        add(f"Sutun {col}", "Sutun", "Gizli", bool(d1.hidden), bool(d2.hidden))

    # Satir yukseklikleri / gizli satirlar
    rows = set(ws1.row_dimensions.keys()) | set(ws2.row_dimensions.keys())
    for row in sorted(rows):
        d1 = ws1.row_dimensions[row]
        d2 = ws2.row_dimensions[row]
        add(f"Satir {row}", "Satir", "Yukseklik", d1.height, d2.height)
        add(f"Satir {row}", "Satir", "Gizli", bool(d1.hidden), bool(d2.hidden))


def compare_cells(name, ws1, ws2, wv1, wv2, diffs, check_format):
    # Sadece dosyada gercekten tanimli hucreleri gez (bos hucreler yaratilmasin diye)
    cells1 = ws1._cells
    cells2 = ws2._cells
    coords = set(cells1.keys()) | set(cells2.keys())

    default_style = None
    for (r, c) in sorted(coords):
        coord = f"{get_column_letter(c)}{r}"
        c1 = cells1.get((r, c))
        c2 = cells2.get((r, c))
        v1 = c1.value if c1 is not None else None
        v2 = c2.value if c2 is not None else None

        # Deger / formul
        if v1 != v2:
            is_formula = any(isinstance(v, str) and v.startswith("=") for v in (v1, v2))
            diffs.append(Diff(name, coord, "Formul" if is_formula else "Deger",
                              "Icerik", v1, v2))
        elif isinstance(v1, str) and v1.startswith("=") and wv1 is not None:
            # Formul ayni ama hesaplanmis (cache) deger farkli olabilir
            cv1 = wv1[coord].value
            cv2 = wv2[coord].value
            if cv1 != cv2:
                diffs.append(Diff(name, coord, "Deger", "Hesaplanmis deger", cv1, cv2))

        # Veri tipi
        if v1 == v2 and c1 is not None and c2 is not None and c1.data_type != c2.data_type:
            diffs.append(Diff(name, coord, "Deger", "Veri tipi", c1.data_type, c2.data_type))

        # Yorumlar
        cm1 = c1.comment.text if c1 is not None and c1.comment else None
        cm2 = c2.comment.text if c2 is not None and c2.comment else None
        if cm1 != cm2:
            diffs.append(Diff(name, coord, "Yorum", "Metin", cm1, cm2))

        # Hyperlink
        h1 = c1.hyperlink.target if c1 is not None and c1.hyperlink else None
        h2 = c2.hyperlink.target if c2 is not None and c2.hyperlink else None
        if h1 != h2:
            diffs.append(Diff(name, coord, "Hyperlink", "Hedef", h1, h2))

        # Format
        if check_format:
            if default_style is None:
                default_style = cell_style(openpyxl.Workbook().active["A1"])
            s1 = cell_style(c1) if c1 is not None else default_style
            s2 = cell_style(c2) if c2 is not None else default_style
            for key in list(s1) + [k for k in s2 if k not in s1]:
                if s1.get(key) != s2.get(key):
                    diffs.append(Diff(name, coord, key[0], key[1], s1.get(key), s2.get(key)))


def compare_workbooks(path1, path2, check_format=True, check_cached=True):
    wb1 = load(path1, data_only=False)
    wb2 = load(path2, data_only=False)
    wbv1 = load(path1, data_only=True) if check_cached else None
    wbv2 = load(path2, data_only=True) if check_cached else None

    diffs = []
    names1, names2 = wb1.sheetnames, wb2.sheetnames

    for n in names1:
        if n not in names2:
            diffs.append(Diff(n, "-", "Sheet", "Varlik", "var", "yok (silinmis)"))
    for n in names2:
        if n not in names1:
            diffs.append(Diff(n, "-", "Sheet", "Varlik", "yok", "var (eklenmis)"))

    common = [n for n in names1 if n in names2]
    order1 = [n for n in names1 if n in common]
    order2 = [n for n in names2 if n in common]
    if order1 != order2:
        diffs.append(Diff("-", "-", "Workbook", "Sheet sirasi",
                          ", ".join(order1), ", ".join(order2)))

    # Tanimli isimler (named ranges)
    try:
        dn1 = {k: v.attr_text for k, v in wb1.defined_names.items()}
        dn2 = {k: v.attr_text for k, v in wb2.defined_names.items()}
        for k in sorted(dn1.keys() | dn2.keys()):
            if dn1.get(k) != dn2.get(k):
                diffs.append(Diff("-", k, "Workbook", "Tanimli isim", dn1.get(k), dn2.get(k)))
    except AttributeError:
        pass

    for n in common:
        ws1, ws2 = wb1[n], wb2[n]
        wv1 = wbv1[n] if wbv1 is not None else None
        wv2 = wbv2[n] if wbv2 is not None else None
        compare_sheet_props(n, ws1, ws2, diffs)
        compare_cells(n, ws1, ws2, wv1, wv2, diffs, check_format)

    return diffs, names1, names2


# --------------------------------------------------------------------------- #
# Rapor
# --------------------------------------------------------------------------- #
HEADERS = ["Sheet", "Konum", "Kategori", "Ozellik", "Dosya 1", "Dosya 2"]


def rows_of(diffs):
    return [[d.sheet, d.location, d.category, d.prop, fmt(d.old), fmt(d.new)] for d in diffs]


def print_report(diffs, path1, path2):
    print(f"Dosya 1: {path1}")
    print(f"Dosya 2: {path2}")
    print("=" * 70)
    if not diffs:
        print("Fark bulunamadi. Dosyalar ayni.")
        return
    current = None
    for d in diffs:
        if d.sheet != current:
            current = d.sheet
            print(f"\n[Sheet: {current}]")
        print(f"  {d.location:<10} {d.category} / {d.prop}: {fmt(d.old)!s}  ->  {fmt(d.new)!s}")
    print("\n" + "=" * 70)
    print(f"Toplam fark: {len(diffs)}")
    summary = {}
    for d in diffs:
        summary[d.category] = summary.get(d.category, 0) + 1
    for cat, cnt in sorted(summary.items(), key=lambda x: -x[1]):
        print(f"  {cat:<20} {cnt}")


def write_csv(diffs, out):
    with open(out, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(HEADERS)
        w.writerows(rows_of(diffs))


def write_xlsx(diffs, out, path1, path2):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ozet"
    ws.append(["Dosya 1", path1])
    ws.append(["Dosya 2", path2])
    ws.append(["Toplam fark", len(diffs)])
    ws.append([])
    ws.append(["Sheet", "Kategori", "Fark sayisi"])
    header_row = ws.max_row
    summary = {}
    for d in diffs:
        summary[(d.sheet, d.category)] = summary.get((d.sheet, d.category), 0) + 1
    for (s, c), n in sorted(summary.items()):
        ws.append([s, c, n])
    bold = Font(bold=True)
    for row in (1, 2, 3):
        ws.cell(row, 1).font = bold
    for cell in ws[header_row]:
        cell.font = bold
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 14

    wd = wb.create_sheet("Farklar")
    wd.append(HEADERS)
    head_fill = PatternFill("solid", fgColor="305496")
    for cell in wd[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
    old_fill = PatternFill("solid", fgColor="FCE4D6")
    new_fill = PatternFill("solid", fgColor="E2EFDA")
    for r in rows_of(diffs):
        wd.append([str(x) for x in r])
        wd.cell(wd.max_row, 5).fill = old_fill
        wd.cell(wd.max_row, 6).fill = new_fill
    for col, width in zip("ABCDEF", (20, 12, 16, 24, 40, 40)):
        wd.column_dimensions[col].width = width
    for row in wd.iter_rows(min_row=2):
        for cell in row[4:6]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wd.freeze_panes = "A2"
    if diffs:
        wd.auto_filter.ref = wd.dimensions
    wb.save(out)


def main(argv=None):
    ap = argparse.ArgumentParser(
        description="Iki Excel dosyasini (tum sheet'ler, degerler ve formatlar) karsilastirir.")
    ap.add_argument("file1", help="Birinci (eski) Excel dosyasi")
    ap.add_argument("file2", help="Ikinci (yeni) Excel dosyasi")
    ap.add_argument("-o", "--output",
                    help="Rapor dosyasi (.xlsx, .csv veya .txt). Verilmezse sadece ekrana yazilir.")
    ap.add_argument("--no-format", action="store_true",
                    help="Format/font farklarini karsilastirma (sadece degerler)")
    ap.add_argument("--no-cached", action="store_true",
                    help="Formullerin hesaplanmis degerlerini karsilastirma")
    ap.add_argument("-q", "--quiet", action="store_true", help="Ekrana detay yazma")
    args = ap.parse_args(argv)

    for p in (args.file1, args.file2):
        if not os.path.isfile(p):
            print(f"Hata: dosya bulunamadi: {p}", file=sys.stderr)
            return 2

    try:
        diffs, _, _ = compare_workbooks(args.file1, args.file2,
                                        check_format=not args.no_format,
                                        check_cached=not args.no_cached)
    except Exception as e:  # noqa: BLE001
        print(f"Hata: dosyalar okunamadi: {e}", file=sys.stderr)
        return 2

    if args.quiet:
        print(f"Toplam fark: {len(diffs)}")
    else:
        print_report(diffs, args.file1, args.file2)

    if args.output:
        ext = os.path.splitext(args.output)[1].lower()
        if ext == ".csv":
            write_csv(diffs, args.output)
        elif ext == ".txt":
            with open(args.output, "w", encoding="utf-8") as fh:
                old = sys.stdout
                sys.stdout = fh
                try:
                    print_report(diffs, args.file1, args.file2)
                finally:
                    sys.stdout = old
        else:
            write_xlsx(diffs, args.output, args.file1, args.file2)
        print(f"\nRapor kaydedildi: {args.output}")

    return 1 if diffs else 0


if __name__ == "__main__":
    sys.exit(main())
